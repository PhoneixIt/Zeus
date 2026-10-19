import openpyxl, sys, io

p = r'C:\Users\Samer\Desktop\AIOS\AIOS_Tool_Stack.xlsx'
wb = openpyxl.load_workbook(p, data_only=False)
out = io.open(r'C:\Users\Samer\Projects\AIOS\_wb_dump.txt', 'w', encoding='utf-8')

out.write('SHEETS: %r\n\n' % wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    out.write('=' * 78 + '\n')
    out.write('SHEET: %s   dims=%s  max_row=%s max_col=%s\n' % (name, ws.dimensions, ws.max_row, ws.max_column))
    out.write('=' * 78 + '\n')
    for row in ws.iter_rows():
        cells = []
        for c in row:
            v = c.value
            if v is None:
                continue
            cells.append('%s=%r' % (c.coordinate, v))
        if cells:
            out.write('  ' + ' | '.join(cells) + '\n')
    out.write('\n')
out.close()
print('dumped')